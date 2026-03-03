import re
from openpyxl import Workbook
from openpyxl.styles import Alignment


def parse_line(line: str) -> dict:
    """Parse a single agenda line into a dict of fields."""
    # 每一行以中文逗號分割，欄位包含「時間」「內容」「講者」
    parts = line.strip().split('，')
    entry: dict = {}
    for part in parts:
        if '：' in part:
            key, val = part.split('：', 1)
            entry[key.strip()] = val.strip()
    return entry


def txt_to_excel(txt_path: str, xlsx_path: str) -> None:
    """Read agenda text file and write to an Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Agenda'

    # 標題列
    ws.append(['時間', '內容', '講者'])

    with open(txt_path, encoding='utf-8') as f:
        for line in f:
            if not line.strip():
                continue
            entry = parse_line(line)

            time = entry.get('時間', '')
            content = entry.get('內容', '')
            speaker = entry.get('講者', '')

            # if there is no speaker (or it's explicitly "無"),
            # we'll still keep the content value and simply merge the two
            # cells so that the content text spans across both columns.
            if not speaker or speaker == '無':
                speaker = ''

            # append row and optionally merge cells
            ws.append([time, content, speaker])
            current_row = ws.max_row
            # center-align the content cell (and merged region if speaker missing)
            cell = ws.cell(row=current_row, column=2)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            if not speaker:
                # merge the "內容" and "講者" columns for this row
                ws.merge_cells(start_row=current_row, start_column=2,
                               end_row=current_row, end_column=3)
                # ensure alignment persists after merge
                merged = ws.cell(row=current_row, column=2)
                merged.alignment = Alignment(horizontal='center', vertical='center')

    wb.save(xlsx_path)
    print(f"已將內容輸出到 {xlsx_path}")


if __name__ == '__main__':
    input_file = '/home/student/exam_test/midterm_agenda.txt'
    output_file = '/home/student/exam_test/midterm_agenda.xlsx'
    txt_to_excel(input_file, output_file)
