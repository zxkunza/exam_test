import re
from openpyxl import Workbook
from openpyxl.styles import Alignment
from typing import List


def parse_line(line: str) -> dict:
    parts = line.strip().split('，')
    entry: dict = {}
    for part in parts:
        if '：' in part:
            key, val = part.split('：', 1)
            entry[key.strip()] = val.strip()
    return entry


def lines_to_excel(lines: List[str], xlsx_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Agenda'

    ws.append(['時間', '內容', '講者'])

    for line in lines:
        if not line.strip():
            continue
        entry = parse_line(line)
        time = entry.get('時間', '')
        content = entry.get('內容', '')
        speaker = entry.get('講者', '')

        if not speaker or speaker == '無':
            speaker = ''

        ws.append([time, content, speaker])
        current_row = ws.max_row

        cell = ws.cell(row=current_row, column=2)
        cell.alignment = Alignment(horizontal='center', vertical='center')

        if not speaker:
            ws.merge_cells(start_row=current_row, start_column=2,
                           end_row=current_row, end_column=3)
            merged = ws.cell(row=current_row, column=2)
            merged.alignment = Alignment(horizontal='center', vertical='center')

    wb.save(xlsx_path)
    print(f"已將內容輸出到 {xlsx_path}")


if __name__ == '__main__':
    text = """\
時間：09:00~09:30，內容：報到，講者：無
時間：09:30~09:40，內容：開場致詞，講者：王教授
時間：09:40~10:05，內容：從 MWC 2026 到 6G：AI-RAN 標準、開源與互通測試的最新進展，講者： 劉教授
時間：10:05~10:30，內容：基於 O-RAN 開放介面的 ISAC 無線感知技術，講者：陳教授
時間：10:30~10:50，內容：Break，講者：無
時間：10:50~11:20，內容：Federated Foundational Models in AI-RAN：Practical and Forward Looking Perspective，講者：教學團隊
時間：11:20~12:00，內容：O-RAN 環境與各模組化功能介紹，講者：教學團隊
時間：12:00~13:30，內容：Lunch，講者：無
時間：13:30~14:00，內容：O-RAN 開源軟體組織簡介，講者：教學團隊
時間：14:00~14:30，內容：O-RAN 實驗環境建置教學，講者：教學團隊
時間：14:30~14:50，內容：Break，講者：無
時間：14:50~15:50，內容：O-RAN xApps 實作建置教學，講者：教學團隊
時間：15:50~16:30，內容：現場討論時間，講者：教學團隊
"""
    lines = text.splitlines()
    lines_to_excel(lines, 'midterm_agenda.xlsx')
